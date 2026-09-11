#!/usr/bin/env python3
"""Open a workbook written by XlsxWriter and report what a real reader sees.

The C++ tests check the bytes; this checks that a spreadsheet library agrees.
Run it after changing anything in src/export/.

    pip install openpyxl
    python3 tools/validate_xlsx.py Worklog_2026_09.xlsx
"""

import sys
import xml.dom.minidom
import zipfile


def check_container(path):
    """Every part must be intact and must be well-formed XML."""
    with zipfile.ZipFile(path) as archive:
        broken = archive.testzip()
        if broken is not None:
            raise SystemExit(f"corrupt entry in the archive: {broken}")

        print(f"{len(archive.namelist())} parts:")
        for name in archive.namelist():
            xml.dom.minidom.parseString(archive.read(name))
            print(f"  ok  {name}  ({len(archive.read(name))} bytes)")


def describe_workbook(path):
    """Print what openpyxl makes of the sheets, styles and links."""
    try:
        import openpyxl
    except ImportError:
        print("\nopenpyxl is not installed; skipping the reader check.")
        print("  pip install openpyxl")
        return

    workbook = openpyxl.load_workbook(path)
    print(f"\nsheets: {workbook.sheetnames}")

    for name in workbook.sheetnames:
        sheet = workbook[name]
        print(f"\n[{name}] {sheet.dimensions}"
              f"  freeze={sheet.freeze_panes}"
              f"  autofilter={sheet.auto_filter.ref}")

        links = [(cell.coordinate, cell.hyperlink.target)
                 for row in sheet.iter_rows() for cell in row if cell.hyperlink]
        if links:
            print(f"  {len(links)} hyperlinks, first: {links[0]}")

        filled = [(cell.coordinate, cell.fill.fgColor.rgb)
                  for row in sheet.iter_rows() for cell in row
                  if cell.fill is not None and cell.fill.patternType == "solid"]
        if filled:
            print(f"  {len(filled)} filled cells, first: {filled[0]}")

        for row in sheet.iter_rows(min_row=1, max_row=3, max_col=6):
            print("  ", [cell.value for cell in row])


def main():
    if len(sys.argv) != 2:
        raise SystemExit(f"usage: {sys.argv[0]} <file.xlsx>")

    path = sys.argv[1]
    check_container(path)
    describe_workbook(path)
    print("\nok")


if __name__ == "__main__":
    main()
